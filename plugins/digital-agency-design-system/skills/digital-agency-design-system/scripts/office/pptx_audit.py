#!/usr/bin/env python3
"""PowerPoint（.pptx）の審査の道具の本体。pptx_audit.sh から呼ばれる（直接呼ばない）。

ファイルは読むだけで書き換えない。基準は references/office-standard.md（S1〜S13）。
- 構造（python-pptx）: S1 タイトル、S2 読み上げ順（タイトルが最初か）、S3 代替テキスト、S6 表の見出し、
  S7 リンクの文字、S8 言語、S10 テーマ外の色・フォント
- 描画（LibreOffice → PDF → pdfminer.six・pypdfium2）: S4 コントラスト、S9 文字の最小サイズ、スライドの画像
出力: result.json（機械向け）、summary.md（審査役が読む日本語の要約）、slide-NN.png
"""
import argparse
import collections
import json
import os
import re
import shutil
import statistics
import subprocess
import sys
import tempfile
from datetime import datetime, timezone

NS_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
NS_P = "http://schemas.openxmlformats.org/presentationml/2006/main"
DECORATIVE_URI = "{C183D7F6-B498-43B3-948B-1728B52AA6E4}"
JA = re.compile(r"[぀-ヿ㐀-鿿]")
VAGUE_LINK = {"こちら", "ここ", "ここをクリック", "詳細", "詳しくはこちら", "リンク", "click here", "here", "link", "more"}
PLACEHOLDER_ALT = re.compile(r"^(図|画像|写真|picture|image|グラフ|chart|graphic)\s*\d*$|\.(png|jpe?g|gif|svg|emf|wmf)$", re.I)
ORDER = {"致命的": 0, "重要": 1, "軽微": 2}
SHOT_WIDTH = 1280


def finding(source, severity, title, slide=None, shape=None, sc=None, level=None, mark=None, examples=None, count=None):
    return {"source": source, "severity": severity, "level": level, "sc": sc or [], "mark": mark,
            "title": title, "slide": slide, "shape": shape, "examples": examples or [], "count": count}


def walk(shapes):
    """グループの中まで含めて、図形を並び順（＝読み上げ順）に返す。"""
    for sh in shapes:
        yield sh
        if sh.shape_type == 6:  # GROUP
            yield from walk(sh.shapes)


def cnvpr(shape):
    found = shape._element.xpath(".//p:cNvPr")
    return found[0] if found else None


def is_decorative(el):
    if el is None:
        return False
    return any(DECORATIVE_URI in (ext.get("uri") or "") for ext in el.iter("{%s}ext" % NS_A))


def label(shape):
    return {"id": shape.shape_id, "name": shape.name}


# ---------------------------------------------------------------- テーマ

def theme_info(prs):
    colors, fonts = set(), set()
    try:
        from lxml import etree
        from pptx.opc.constants import RELATIONSHIP_TYPE as RT
        for master in prs.slide_masters:
            theme = etree.fromstring(master.part.part_related_by(RT.THEME).blob)
            for el in theme.iter("{%s}srgbClr" % NS_A):
                colors.add(el.get("val").upper())
            for el in theme.iter("{%s}sysClr" % NS_A):
                if el.get("lastClr"):
                    colors.add(el.get("lastClr").upper())
            for tag in ("latin", "ea"):
                for el in theme.iter("{%s}%s" % (NS_A, tag)):
                    if el.get("typeface"):
                        fonts.add(el.get("typeface"))
            # マスターとレイアウトに直接書かれた色・フォントもテンプレートの一部として認める
            parts = [master] + list(master.slide_layouts)
            for part_owner in parts:
                for el in part_owner._element.iter("{%s}srgbClr" % NS_A):
                    colors.add(el.get("val").upper())
                for tag in ("latin", "ea"):
                    for el in part_owner._element.iter("{%s}%s" % (NS_A, tag)):
                        if el.get("typeface"):
                            fonts.add(el.get("typeface"))
    except Exception:
        pass
    colors |= {"FFFFFF", "000000"}
    return colors, fonts


def norm(t):
    return re.sub(r"\s+", "", t or "")


def template_texts(prs):
    """マスターとレイアウトに書かれた文字（フッターなど）。テンプレート由来として作成者の指摘から外す。"""
    texts = set()
    for master in prs.slide_masters:
        for owner in [master] + list(master.slide_layouts):
            for sh in walk(owner.shapes):
                if sh.is_placeholder or not sh.has_text_frame:
                    continue
                t = norm(sh.text_frame.text)
                if len(t) >= 2:
                    texts.add(t)
    return texts


def palette_colors():
    """references/office-palettes.json の色（slide-techvan など、テーマの色と同じに扱う色）。"""
    p = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..", "references", "office-palettes.json")
    try:
        data = json.load(open(p, encoding="utf-8"))
    except Exception:
        return {}
    return {name: {c.upper() for c in v.get("colors", [])} for name, v in data.items() if isinstance(v, dict)}


def protect_colors(target):
    """DADS-PROTECT.md の「ブランド色」に書かれた色（資料のフォルダから上へ探す）。"""
    d = os.path.dirname(os.path.abspath(target))
    while True:
        p = os.path.join(d, "DADS-PROTECT.md")
        if os.path.exists(p):
            text = open(p, encoding="utf-8").read()
            m = re.search(r"^##\s*ブランド色\s*$(.*?)(^##\s|\Z)", text, re.M | re.S)
            return p, {c.upper() for c in re.findall(r"#([0-9A-Fa-f]{6})\b", m.group(1))} if m else set()
        if os.path.exists(os.path.join(d, ".git")) or os.path.dirname(d) == d:
            return None, set()
        d = os.path.dirname(d)


def default_lang(prs):
    el = prs.part._element.find(".//{%s}defaultTextStyle" % NS_P)
    if el is not None:
        rpr = el.find(".//{%s}defRPr" % NS_A)
        if rpr is not None and rpr.get("lang"):
            return rpr.get("lang")
    return None


# ---------------------------------------------------------------- 構造の点検

def structure(prs, brand_colors):
    out = []
    theme_colors, theme_fonts = theme_info(prs)
    theme_colors = theme_colors | brand_colors
    deflang = default_lang(prs)
    titles = collections.defaultdict(list)
    lang_bad, lang_missing = [], []
    color_bad = collections.Counter()
    color_example = {}
    font_bad = collections.Counter()
    font_example = {}
    slides_info = []

    for no, slide in enumerate(prs.slides, start=1):
        title_shape = slide.shapes.title
        title = (title_shape.text_frame.text.strip() if title_shape is not None and title_shape.has_text_frame else "")
        images = []
        slides_info.append({"slide": no, "title": title, "hidden": slide._element.get("show") == "0", "images": images})

        # S1 タイトル
        if title_shape is None:
            out.append(finding("S1:no-title", "重要", "スライドにタイトルのプレースホルダが無い", no, sc=["2.4.6"], level="AA"))
        elif not title:
            out.append(finding("S1:empty-title", "重要", "スライドのタイトルが空", no, label(title_shape), sc=["2.4.6"], level="AA"))
        else:
            titles[title].append(no)

        # S2 読み上げ順：タイトルが最初の中身か
        content = [sh for sh in slide.shapes if (sh.has_text_frame and sh.text_frame.text.strip()) or sh.shape_type in (13, 6) or getattr(sh, "has_table", False) or getattr(sh, "has_chart", False)]
        if title_shape is not None and title and content and content[0].shape_id != title_shape.shape_id:
            out.append(finding("S2:title-not-first", "致命的", "タイトルより先に読み上げられる図形がある（タイトルを最初にする）",
                               no, label(content[0]), sc=["1.3.2"], level="A", examples=[f"最初: {content[0].name}"]))

        for sh in walk(slide.shapes):
            el = cnvpr(sh)
            is_pic = sh.shape_type == 13 or (sh.shape_type == 14 and sh._element.find(".//{%s}blip" % NS_A) is not None)
            is_graphic = getattr(sh, "has_chart", False) or sh._element.find(".//{http://schemas.openxmlformats.org/drawingml/2006/diagram}relIds") is not None
            if is_pic or is_graphic:
                images.append({"id": sh.shape_id, "name": sh.name, "kind": "graphic" if is_graphic else "picture",
                               "alt": (el.get("descr") or "") if el is not None else "", "decorative": is_decorative(el)})
            # S3 代替テキスト
            if (is_pic or is_graphic) and not is_decorative(el):
                descr = (el.get("descr") or "").strip() if el is not None else ""
                kind = "グラフ・図" if is_graphic else "画像"
                if not descr:
                    out.append(finding("S3:no-alt", "致命的", f"{kind}に代替テキストが無い（飾りなら「装飾用」にする）", no, label(sh), sc=["1.1.1"], level="A"))
                elif PLACEHOLDER_ALT.search(descr) or "自動生成" in descr:
                    out.append(finding("S3:placeholder-alt", "致命的", f"{kind}の代替テキストが仮の名前・自動生成のまま", no, label(sh), sc=["1.1.1"], level="A", examples=[descr[:60]]))
            # S6 表の見出し
            if getattr(sh, "has_table", False) and not sh.table.first_row:
                out.append(finding("S6:no-header", "致命的", "表の1行目が見出しに設定されていない", no, label(sh), sc=["1.3.1"], level="A"))
            if not sh.has_text_frame:
                continue
            for para in sh.text_frame.paragraphs:
                for run in para.runs:
                    text = run.text.strip()
                    if not text:
                        continue
                    # S7 リンクの文字
                    try:
                        addr = run.hyperlink.address
                    except Exception:
                        addr = None
                    if addr and (text.lower() in VAGUE_LINK or text == addr or text.startswith("http")):
                        out.append(finding("S7:vague-link", "致命的", "リンクの文字だけでは行き先が分からない", no, label(sh), sc=["2.4.4"], level="A", examples=[text[:40]]))
                    rpr = run._r.find("{%s}rPr" % NS_A)
                    # S8 言語
                    if JA.search(text):
                        lang = rpr.get("lang") if rpr is not None else None
                        if lang and not lang.lower().startswith("ja"):
                            lang_bad.append((no, sh.name, lang, text[:20]))
                        elif not lang and deflang and not deflang.lower().startswith("ja"):
                            lang_missing.append((no, sh.name, deflang, text[:20]))
                    # S10 テーマ外の色・フォント
                    if rpr is not None:
                        for c in rpr.iter("{%s}srgbClr" % NS_A):
                            v = c.get("val").upper()
                            if v not in theme_colors:
                                color_bad[v] += 1
                                color_example.setdefault(v, f"スライド{no}「{text[:15]}」")
                        for tag in ("latin", "ea"):
                            f = rpr.find("{%s}%s" % (NS_A, tag))
                            if f is not None and f.get("typeface") and not f.get("typeface").startswith("+") and f.get("typeface") not in theme_fonts:
                                font_bad[f.get("typeface")] += 1
                                font_example.setdefault(f.get("typeface"), f"スライド{no}「{text[:15]}」")
            # 図形の塗りのテーマ外の色
            sppr = sh._element.find("{%s}spPr" % NS_P)
            if sppr is not None:
                fill = sppr.find("{%s}solidFill" % NS_A)
                if fill is not None:
                    c = fill.find("{%s}srgbClr" % NS_A)
                    if c is not None and c.get("val").upper() not in theme_colors:
                        v = c.get("val").upper()
                        color_bad[v] += 1
                        color_example.setdefault(v, f"スライド{no}の図形「{sh.name}」の塗り")

    for t, nos in titles.items():
        if len(nos) > 1:
            out.append(finding("S1:duplicate-title", "重要", "同じタイトルのスライドがある（続きは「（2/3）」などで区別する）",
                               nos[0], sc=["2.4.6"], level="AA", examples=[f"「{t[:30]}」: スライド {', '.join(map(str, nos))}"]))
    if lang_bad:
        out.append(finding("S8:lang", "重要", "日本語の文字に日本語以外の言語設定が付いている", lang_bad[0][0], sc=["3.1.2"], level="AA",
                           count=len(lang_bad), examples=[f"スライド{n}「{t}」（{l}）" for n, _, l, t in lang_bad[:3]]))
    if lang_missing:
        out.append(finding("S8:default-lang", "重要", f"既定の言語が日本語以外（{lang_missing[0][2]}）で、日本語の文字に言語設定が無い", lang_missing[0][0],
                           sc=["3.1.2"], level="AA", count=len(lang_missing), examples=[f"スライド{n}「{t}」" for n, _, _, t in lang_missing[:3]]))
    if color_bad:
        out.append(finding("S10:color", "軽微", f"テンプレートのテーマにない色を直接指定している（{len(color_bad)} 色。使い続けるなら DADS-PROTECT.md の「ブランド色」に書く）", mark="基準", count=len(color_bad),
                           examples=[f"#{c} を {n} か所（例：{color_example[c]}）" for c, n in color_bad.most_common(6)]))
    if font_bad:
        out.append(finding("S10:font", "軽微", f"テンプレートのテーマにないフォントを直接指定している（{len(font_bad)} 種類）", mark="基準", count=len(font_bad),
                           examples=[f"{f} を {n} か所（例：{font_example[f]}）" for f, n in font_bad.most_common(5)]))
    return out, slides_info, {"colors": sorted(theme_colors), "fonts": sorted(theme_fonts), "defaultLang": deflang}


# ---------------------------------------------------------------- 描画の点検

def to_pdf(soffice, path, workdir):
    profile = tempfile.mkdtemp(prefix="lo-profile-")
    try:
        subprocess.run([soffice, f"-env:UserInstallation=file://{profile}", "--headless", "--convert-to", "pdf",
                        "--outdir", workdir, path], check=True, capture_output=True, timeout=300)
    finally:
        shutil.rmtree(profile, ignore_errors=True)
    pdf = os.path.join(workdir, os.path.splitext(os.path.basename(path))[0] + ".pdf")
    return pdf if os.path.exists(pdf) else None


def rgb_of(color):
    """pdfminer の色（グレー・RGB・CMYK の 0〜1）を 0〜255 の RGB にする。"""
    if color is None:
        return (0, 0, 0)
    if isinstance(color, (int, float)):
        color = (color,)
    c = [float(x) for x in color]
    if len(c) == 1:
        v = round(c[0] * 255)
        return (v, v, v)
    if len(c) == 3:
        return tuple(round(x * 255) for x in c)
    if len(c) == 4:
        k = c[3]
        return tuple(round(255 * (1 - x) * (1 - k)) for x in c[:3])
    return (0, 0, 0)


def lum(rgb):
    def ch(v):
        v /= 255
        return v / 12.92 if v <= 0.03928 else ((v + 0.055) / 1.055) ** 2.4
    r, g, b = (ch(x) for x in rgb)
    return 0.2126 * r + 0.7152 * g + 0.0722 * b


def ratio(a, b):
    la, lb = lum(a), lum(b)
    return (max(la, lb) + 0.05) / (min(la, lb) + 0.05)


def background(img, box, text_rgb):
    """文字の範囲の画素のうち、文字色から最も遠い色を多く含む色を背景とみなす。"""
    x0, y0, x1, y1 = box
    crop = img.crop((max(0, x0), max(0, y0), min(img.width, x1), min(img.height, y1))).convert("RGB")
    if crop.width < 2 or crop.height < 2:
        return None, 1.0
    pixels = crop.get_flattened_data() if hasattr(crop, "get_flattened_data") else crop.getdata()
    q = collections.Counter((r // 8 * 8, g // 8 * 8, b // 8 * 8) for r, g, b in pixels)
    total = sum(q.values())
    cands = [(n, c) for c, n in q.most_common(6) if sum(abs(a - b) for a, b in zip(c, text_rgb)) > 60]
    if not cands:
        return None, 1.0
    n, c = cands[0]
    return c, n / total


def render(soffice, path, outdir):
    from pdfminer.high_level import extract_pages
    from pdfminer.layout import LAParams, LTChar, LTTextContainer, LTTextLine
    import pypdfium2 as pdfium

    work = tempfile.mkdtemp(prefix="pptx-audit-")
    try:
        pdf = to_pdf(soffice, path, work)
        if not pdf:
            return None
        doc = pdfium.PdfDocument(pdf)
        pages = []
        for i, page in enumerate(extract_pages(pdf, laparams=LAParams()), start=1):
            pw, ph = page.width, page.height
            scale = SHOT_WIDTH / pw
            img = doc[i - 1].render(scale=scale).to_pil()
            shot = f"slide-{i:02d}.png"
            img.save(os.path.join(outdir, shot))
            lines = []

            def collect(obj):
                for child in obj:
                    if isinstance(child, LTTextLine):
                        chars = [c for c in child if isinstance(c, LTChar) and c.get_text().strip()]
                        if not chars:
                            continue
                        size = round(statistics.median(c.size for c in chars) * 2) / 2
                        bold = sum("bold" in c.fontname.lower() or "w6" in c.fontname.lower() or "heavy" in c.fontname.lower() for c in chars) > len(chars) / 2
                        color = collections.Counter(rgb_of(c.graphicstate.ncolor) for c in chars).most_common(1)[0][0]
                        x0, y0, x1, y1 = child.bbox
                        box = (int(x0 * scale), int((ph - y1) * scale), int(x1 * scale), int((ph - y0) * scale))
                        bg, share = background(img, box, color)
                        lines.append({"text": child.get_text().strip()[:60], "size": size, "bold": bold,
                                      "color": "#%02X%02X%02X" % color,
                                      "background": ("#%02X%02X%02X" % bg) if bg else None,
                                      "contrast": round(ratio(color, bg), 2) if bg else None,
                                      "busyBackground": bool(bg) and share < 0.5})
                    elif isinstance(child, LTTextContainer) or hasattr(child, "__iter__"):
                        try:
                            collect(child)
                        except TypeError:
                            pass

            collect(page)
            pages.append({"slide": i, "screenshot": shot, "lines": lines})
        return pages
    finally:
        shutil.rmtree(work, ignore_errors=True)


def render_findings(pages, tmpl):
    out = []
    small, mid, low, from_template = [], [], [], []
    for p in pages:
        for ln in p["lines"]:
            t = norm(ln["text"])
            if t and any(t == x or (len(t) >= 6 and t in x) or (len(x) >= 6 and x in t) for x in tmpl):
                ln["template"] = True
                from_template.append((p["slide"], ln))
                continue
            if ln["size"] < 12:
                small.append((p["slide"], ln))
            elif ln["size"] < 18:
                mid.append((p["slide"], ln))
            if ln["contrast"] is not None:
                large = ln["size"] >= 18 or (ln["bold"] and ln["size"] >= 14)
                need = 3.0 if large else 4.5
                if ln["contrast"] < need:
                    low.append((p["slide"], ln, need))
    if low:
        out.append(finding("S4:contrast", "重要", "文字と背景のコントラストが足りない", low[0][0], sc=["1.4.3"], level="AA", count=len(low),
                           examples=[f"スライド{s}「{l['text'][:15]}」{l['color']} / {l['background']} = {l['contrast']}:1（必要 {n}:1）{'（背景が画像のため推定）' if l['busyBackground'] else ''}" for s, l, n in low[:5]]))
    if small:
        out.append(finding("S9:tiny", "軽微", "12pt 未満の文字がある（補助の文字も 12pt 以上）", small[0][0], mark="基準", count=len(small),
                           examples=[f"スライド{s}「{l['text'][:15]}」{l['size']}pt" for s, l in small[:5]]))
    tmpl_examples = sorted({f"「{l['text'][:20]}」{l['size']}pt" for _, l in from_template})
    return out, {"under12": len(small), "from12to18": len(mid), "template": len(from_template), "templateExamples": tmpl_examples[:5],
                 "from12to18Examples": [f"スライド{s}「{l['text'][:15]}」{l['size']}pt" for s, l in mid[:8]]}


# ---------------------------------------------------------------- まとめ

def summary(r):
    L = []
    n = lambda s: sum(1 for f in r["findings"] if f["severity"] == s)
    L.append(f"# 審査の道具の記録（PowerPoint）：{os.path.basename(r['target'])}")
    L.append("")
    L.append(f"- 対象: {r['target']}")
    L.append(f"- 日時: {r['checkedAt']}")
    L.append(f"- スライド数: {len(r['slides'])}")
    L.append(f"- 基準: references/office-standard.md（スライド編・初版）")
    L.append(f"- 候補の数: 致命的 {n('致命的')}／重要 {n('重要')}／軽微 {n('軽微')}（基準の印 {sum(1 for f in r['findings'] if f['mark'] == '基準')}）")
    if not r["rendered"]:
        L.append(f"- **描画の点検は未点検**: {r['renderNote']}")
    L.append("")
    L.append("重大度は目安です（WCAG のレベル A＝致命的、AA＝重要、基準だけに当たるもの＝軽微）。審査役がスライドの画像と中身で確かめてから指摘にします。")
    L.append("")
    L.append("## 問題の候補")
    L.append("")
    if not r["findings"]:
        L.append("（道具では見つかりませんでした。目で確かめる観点は残っています）")
    for f in r["findings"]:
        meta = [f"WCAG {f['level']}" if f["level"] else None, f"達成基準 {', '.join(f['sc'])}" if f["sc"] else None,
                "基準" if f["mark"] else None, f"{f['count']} 件" if f["count"] else None,
                f"スライド {f['slide']}" if f["slide"] else None,
                f"図形 {f['shape']['name']}（id {f['shape']['id']}）" if f["shape"] else None]
        L.append(f"- **{f['severity']}** {f['title']}（{'・'.join(m for m in meta if m) or f['source']}）")
        for e in f["examples"]:
            L.append(f"  - {e}")
    L.append("")
    L.append("## 確かめたこと")
    L.append("")
    if r["theme"].get("palettes"):
        L.append(f"- テーマの色と同じに扱った色の一覧: {', '.join(r['theme']['palettes'])}（references/office-palettes.json）")
    if r["theme"].get("protectFile"):
        L.append(f"- 守るものの宣言: {r['theme']['protectFile']}（ブランド色 {len(r['theme']['brandColors'])} 色をテーマの色と同じに扱った）")
    L.append(f"- テーマの色 {len(r['theme']['colors'])} 色、テーマのフォント: {', '.join(f for f in r['theme']['fonts'] if not f.startswith('+')) or '不明'}、既定の言語: {r['theme']['defaultLang'] or '不明'}")
    if r["rendered"]:
        s = r["sizes"]
        L.append(f"- 文字サイズ: 12pt 未満 {s['under12']} 行、12〜18pt {s['from12to18']} 行（本文なら 18pt 以上にする。注記・出典なら 12pt 以上でよい）")
        for e in s["from12to18Examples"]:
            L.append(f"  - {e}")
        if s["template"]:
            L.append(f"- テンプレート（マスター・レイアウト）由来の文字 {s['template']} 行は、作成者の指摘から外した（直すならテンプレートの持ち主に確認する）: {'、'.join(s['templateExamples'])}")
    L.append("")
    L.append("## スライドの一覧")
    L.append("")
    for sl in r["slides"]:
        shot = f"（{sl['screenshot']}）" if sl.get("screenshot") else ""
        L.append(f"- スライド {sl['slide']}: {sl['title'] or '（タイトルなし）'}{'（非表示）' if sl['hidden'] else ''}{shot}")
    L.append("")
    L.append("## 道具では確かめられないこと（画像と中身を見て判断する）")
    L.append("")
    L.append("- S5 色だけで伝えていないか、S11 文字量と行間、S12 はみ出しと重なり、S13 動き")
    L.append("- 代替テキストの中身が画像の役割に合っているか、本文の読み上げ順が見た目の順か")
    L.append("- 画像の上の文字のコントラスト（背景が画像のときは推定値）")
    return "\n".join(L) + "\n"


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--pydir", required=True)
    ap.add_argument("--target", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--soffice", default="")
    a = ap.parse_args()
    sys.path.insert(0, a.pydir)
    from pptx import Presentation

    try:
        prs = Presentation(a.target)
    except Exception as e:
        print(f"PowerPoint のファイルとして開けませんでした: {e}", file=sys.stderr)
        sys.exit(4)

    protect_path, brand = protect_colors(a.target)
    palettes = palette_colors()
    allowed = set(brand)
    for cols in palettes.values():
        allowed |= cols
    findings, slides, theme = structure(prs, allowed)
    theme["palettes"] = sorted(palettes)
    theme["protectFile"] = protect_path
    theme["brandColors"] = sorted(brand)
    r = {"target": os.path.abspath(a.target), "checkedAt": datetime.now(timezone.utc).isoformat(timespec="seconds"),
         "standard": "office-standard.md スライド編・初版", "theme": theme, "slides": slides, "rendered": False, "renderNote": ""}
    if a.soffice:
        try:
            pages = render(a.soffice, a.target, a.out)
        except Exception as e:
            pages = None
            r["renderNote"] = f"PDF への変換・読み取りに失敗した（{str(e)[:120]}）"
        if pages:
            r["rendered"] = True
            if len(pages) != len(slides):
                r["renderNote"] = f"PDF のページ数（{len(pages)}）とスライド数（{len(slides)}）が違う。非表示のスライドは画像にならない"
            visible = [s for s in slides if not s["hidden"]] if len(pages) != len(slides) else slides
            for s, p in zip(visible, pages):
                s["screenshot"] = p["screenshot"]
                p["slide"] = s["slide"]
            rf, sizes = render_findings(pages, template_texts(prs))
            findings += rf
            r["sizes"] = sizes
            r["pages"] = pages
        elif not r["renderNote"]:
            r["renderNote"] = "PDF への変換に失敗した"
    else:
        r["renderNote"] = "LibreOffice（soffice）が無いため、文字サイズ・コントラスト・画像を確かめていない"
    r["findings"] = sorted(findings, key=lambda f: (ORDER[f["severity"]], f["slide"] or 0))
    with open(os.path.join(a.out, "result.json"), "w", encoding="utf-8") as f:
        json.dump(r, f, ensure_ascii=False, indent=2)
    with open(os.path.join(a.out, "summary.md"), "w", encoding="utf-8") as f:
        f.write(summary(r))
    print(os.path.join(a.out, "summary.md"))


if __name__ == "__main__":
    main()
