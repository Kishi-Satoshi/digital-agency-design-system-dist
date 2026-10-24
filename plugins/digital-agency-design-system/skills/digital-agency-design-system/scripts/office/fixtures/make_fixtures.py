#!/usr/bin/env python3
"""審査の道具の確認用 PowerPoint を作る（bad.pptx：わざと問題を入れる／good.pptx：問題なし）。

使い方: PYTHONPATH=<python-pptx の場所> python3 make_fixtures.py
"""
import io
import os

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
NS_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def png():
    buf = io.BytesIO()
    Image.new("RGB", (200, 120), (0, 80, 160)).save(buf, "PNG")
    buf.seek(0)
    return buf


def set_lang(run, lang):
    run._r.get_or_add_rPr().set("lang", lang)


def text(shape, s, size, lang="ja-JP", rgb=None):
    tf = shape.text_frame
    tf.text = ""
    run = tf.paragraphs[0].add_run()
    run.text = s
    run.font.size = Pt(size)
    set_lang(run, lang)
    if rgb:
        run.font.color.rgb = RGBColor.from_string(rgb)
    return run


def bad():
    prs = Presentation()
    # 1枚目：タイトルより先に読まれる図形、薄い文字、小さい文字、テーマ外の色とフォント、言語の誤り、あいまいなリンク
    s = prs.slides.add_slide(prs.slide_layouts[5])  # タイトルのみ
    text(s.shapes.title, "申請の流れ", 32)
    box = s.shapes.add_textbox(Inches(0.5), Inches(1.6), Inches(9), Inches(0.8))
    text(box, "薄くて読みにくい説明文です", 20, rgb="BBBBBB")
    s.shapes._spTree.remove(box._element)
    s.shapes._spTree.insert(2, box._element)  # タイトルより前に置く（読み上げ順の誤り）
    small = s.shapes.add_textbox(Inches(0.5), Inches(2.6), Inches(9), Inches(0.5))
    text(small, "出典：とても小さい文字の注記", 9)
    loud = s.shapes.add_textbox(Inches(0.5), Inches(3.3), Inches(9), Inches(0.6))
    r = text(loud, "テーマにない色とフォント", 20, lang="en-US", rgb="FF00AA")
    r.font.name = "Comic Sans MS"
    r._r.get_or_add_rPr().append(r._r.get_or_add_rPr().makeelement(NS_A + "ea", {"typeface": "Comic Sans MS"}))
    link = s.shapes.add_textbox(Inches(0.5), Inches(4.2), Inches(9), Inches(0.6))
    lr = text(link, "こちら", 20)
    lr.hyperlink.address = "https://example.invalid/guide"
    s.shapes.add_picture(png(), Inches(0.5), Inches(5.0), Inches(2), Inches(1.2))  # 代替テキストなし
    # 2枚目：タイトルなし、見出しの無い表
    s2 = prs.slides.add_slide(prs.slide_layouts[6])  # 白紙
    tbl = s2.shapes.add_table(3, 2, Inches(0.5), Inches(1), Inches(6), Inches(1.5)).table
    tbl.first_row = False
    for i, row in enumerate([("項目", "内容"), ("期限", "毎月末"), ("窓口", "総務")]):
        for j, v in enumerate(row):
            text(tbl.cell(i, j), v, 18)
    # 3枚目と4枚目：同じタイトル
    for _ in range(2):
        s3 = prs.slides.add_slide(prs.slide_layouts[5])
        text(s3.shapes.title, "注意事項", 32)
    prs.save(os.path.join(HERE, "bad.pptx"))


def good():
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[5])
    text(s.shapes.title, "申請の流れ", 32)
    body = s.shapes.add_textbox(Inches(0.5), Inches(1.6), Inches(9), Inches(0.8))
    text(body, "申請は毎月末までに総務へ出します", 20)
    pic = s.shapes.add_picture(png(), Inches(0.5), Inches(2.8), Inches(2), Inches(1.2))
    pic._element.xpath(".//p:cNvPr")[0].set("descr", "申請から承認までの3つの段階を示す図")
    tbl = s.shapes.add_table(2, 2, Inches(3), Inches(2.8), Inches(6), Inches(1)).table
    tbl.first_row = True
    for i, row in enumerate([("項目", "内容"), ("期限", "毎月末")]):
        for j, v in enumerate(row):
            text(tbl.cell(i, j), v, 18, rgb="FFFFFF" if i == 0 else None)
    s2 = prs.slides.add_slide(prs.slide_layouts[5])
    text(s2.shapes.title, "問い合わせ先", 32)
    b2 = s2.shapes.add_textbox(Inches(0.5), Inches(1.6), Inches(9), Inches(0.8))
    lr = text(b2, "申請の手引き（社内ポータル）", 20)
    lr.hyperlink.address = "https://example.invalid/guide"
    prs.save(os.path.join(HERE, "good.pptx"))


if __name__ == "__main__":
    bad()
    good()
    print("bad.pptx と good.pptx を作りました")
