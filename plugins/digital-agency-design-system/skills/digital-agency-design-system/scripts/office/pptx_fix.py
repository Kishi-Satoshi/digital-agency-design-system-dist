#!/usr/bin/env python3
"""PowerPoint（.pptx）の指摘を直す道具。元のファイルは変えず、別名（既定: <名前>.dads-fixed.pptx）に保存する。

スライドの中身だけを書き換え、スライドマスター・レイアウト・テーマには手を付けない。
直す内容は JSON の配列で渡す（直す役が審査の記録 result.json の slide・shape.id を見て作る）。

使い方:
  python3 pptx_fix.py <資料.pptx> <直す内容.json> [--out <出力.pptx>] [--pydir <python-pptx の場所>]

直す内容の例:
  [
    {"op": "alt", "slide": 1, "shape": 7, "text": "申請から承認までの3段階を示す図"},
    {"op": "decorative", "slide": 1, "shape": 8},
    {"op": "title-first", "slide": 1},
    {"op": "title", "slide": 4, "text": "注意事項（2/2）"},
    {"op": "table-header", "slide": 2, "shape": 2},
    {"op": "font-size", "slide": 1, "shape": 4, "pt": 12},
    {"op": "color", "slide": 1, "shape": 3, "theme": "tx1"},
    {"op": "color", "slide": 1, "shape": 3, "rgb": "333333"},
    {"op": "theme-font", "slide": 1, "shape": 5},
    {"op": "lang", "lang": "ja-JP"},
    {"op": "color", "slide": 1, "find": "薄くて読みにくい", "theme": "tx1"},
    {"op": "replace-text", "slide": 1, "shape": 6, "old": "こちら", "new": "申請の手引き（社内ポータル）"}
  ]
図形は shape（審査の記録の図形 id）か find（図形の中の文字の一部。1つに絞れるもの）で指定する。
op の意味:
  alt           画像・図・グラフに代替テキストを付ける（S3）
  decorative    飾りの画像を「装飾用」にする（S3）
  title-first   タイトルを読み上げ順の最初にする（S2）
  title         空のタイトル、重複したタイトルを書き換える（S1）。タイトルのプレースホルダが無いスライドは直せない（レイアウトの変更が要るので確認する）
  table-header  表の1行目を見出しにする（S6）
  font-size     図形の中の文字を pt 以上にする（指定より小さい文字と、サイズの指定が無い文字だけ）（S9）
  color         図形の中の文字の色を変える。theme（tx1・tx2・bg1・bg2・accent1〜6）を優先し、rgb は最後の手段（S4）
  theme-font    図形の中の文字に直接指定したフォントを外し、テーマのフォントに戻す（S10）
  lang          日本語の文字に言語設定 lang を付ける（S8）
  replace-text  図形の中の文字を置き換える（S7 リンクの文字など）
終了コード: 0 = 保存した, 1 = 直せない指示があった（何も保存しない）, 2 = 使い方の誤り
"""
import argparse
import json
import os
import re
import sys

NS_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
DECORATIVE_URI = "{C183D7F6-B498-43B3-948B-1728B52AA6E4}"
NS_ADEC = "http://schemas.microsoft.com/office/drawing/2017/decorative"
JA = re.compile(r"[぀-ヿ㐀-鿿]")
THEME = {"tx1": "TEXT_1", "tx2": "TEXT_2", "bg1": "BACKGROUND_1", "bg2": "BACKGROUND_2",
         **{f"accent{i}": f"ACCENT_{i}" for i in range(1, 7)}}


class FixError(Exception):
    pass


def walk(shapes):
    for sh in shapes:
        yield sh
        if sh.shape_type == 6:
            yield from walk(sh.shapes)


def find_shape(prs, op):
    """op の slide と、shape（図形の id）か find（図形の中の文字の一部）で図形を探す。"""
    slide_no = op["slide"]
    if not 1 <= slide_no <= len(prs.slides):
        raise FixError(f"スライド {slide_no} がありません")
    slide = prs.slides[slide_no - 1]
    hits = []
    for sh in walk(slide.shapes):
        if "shape" in op and sh.shape_id == op["shape"]:
            return slide, sh
        if "find" in op and op["find"] and op["find"] in "".join(r.text for r in runs(sh)):
            hits.append(sh)
    if "find" in op:
        if len(hits) == 1:
            return slide, hits[0]
        raise FixError(f"スライド {slide_no} で「{op['find']}」を含む図形が {len(hits)} 個あります（1つに絞れる文字か、shape の id を指定する）")
    raise FixError(f"スライド {slide_no} に図形 id {op.get('shape')} がありません")


def cnvpr(shape):
    found = shape._element.xpath(".//p:cNvPr")
    if not found:
        raise FixError(f"図形 {shape.name} に代替テキストの欄がありません")
    return found[0]


def runs(shape):
    if shape.has_text_frame:
        for para in shape.text_frame.paragraphs:
            yield from para.runs
    if getattr(shape, "has_table", False):
        for row in shape.table.rows:
            for cell in row.cells:
                for para in cell.text_frame.paragraphs:
                    yield from para.runs


def apply(prs, op):
    from lxml import etree
    from pptx.dml.color import RGBColor
    from pptx.enum.dml import MSO_THEME_COLOR
    from pptx.util import Pt

    kind = op.get("op")
    if kind == "alt":
        _, sh = find_shape(prs, op)
        text = (op.get("text") or "").strip()
        if not text:
            raise FixError("alt には text が必要です（飾りなら decorative を使う）")
        cnvpr(sh).set("descr", text)
        return f"スライド{op['slide']}「{sh.name}」に代替テキスト「{text[:30]}」"
    if kind == "decorative":
        _, sh = find_shape(prs, op)
        el = cnvpr(sh)
        el.set("descr", "")
        ext_lst = el.find("{%s}extLst" % NS_A)
        if ext_lst is None:
            ext_lst = etree.SubElement(el, "{%s}extLst" % NS_A)
        if not any(DECORATIVE_URI in (e.get("uri") or "") for e in ext_lst):
            ext = etree.SubElement(ext_lst, "{%s}ext" % NS_A, uri=DECORATIVE_URI)
            etree.SubElement(ext, "{%s}decorative" % NS_ADEC, nsmap={"adec": NS_ADEC}, val="1")
        return f"スライド{op['slide']}「{sh.name}」を装飾用に"
    if kind == "title-first":
        slide = prs.slides[op["slide"] - 1]
        title = slide.shapes.title
        if title is None:
            raise FixError(f"スライド {op['slide']} にタイトルがありません")
        tree = slide.shapes._spTree
        tree.remove(title._element)
        tree.insert(2, title._element)  # nvGrpSpPr・grpSpPr の次（最初に読まれる位置）
        return f"スライド{op['slide']}のタイトルを読み上げ順の最初に"
    if kind == "title":
        slide = prs.slides[op["slide"] - 1]
        if slide.shapes.title is None:
            raise FixError(f"スライド {op['slide']} にタイトルのプレースホルダがありません（レイアウトの変更が要るので、利用者に確認する）")
        slide.shapes.title.text_frame.text = op["text"]
        if JA.search(op["text"]):
            for r in runs(slide.shapes.title):
                r._r.get_or_add_rPr().set("lang", op.get("lang", "ja-JP"))
        return f"スライド{op['slide']}のタイトルを「{op['text'][:30]}」に"
    if kind == "table-header":
        _, sh = find_shape(prs, op)
        if not getattr(sh, "has_table", False):
            raise FixError(f"図形 {sh.name} は表ではありません")
        sh.table.first_row = True
        return f"スライド{op['slide']}の表「{sh.name}」の1行目を見出しに"
    if kind == "font-size":
        _, sh = find_shape(prs, op)
        pt = float(op["pt"])
        n = 0
        for r in runs(sh):
            if r.font.size is None or r.font.size.pt < pt:
                r.font.size = Pt(pt)
                n += 1
        return f"スライド{op['slide']}「{sh.name}」の文字 {n} か所を {pt:g}pt 以上に"
    if kind == "color":
        _, sh = find_shape(prs, op)
        n = 0
        for r in runs(sh):
            if op.get("theme"):
                if op["theme"] not in THEME:
                    raise FixError(f"theme は {', '.join(THEME)} のどれかにしてください")
                r.font.color.theme_color = getattr(MSO_THEME_COLOR, THEME[op["theme"]])
            elif op.get("rgb"):
                r.font.color.rgb = RGBColor.from_string(op["rgb"].lstrip("#").upper())
            else:
                raise FixError("color には theme か rgb が必要です")
            n += 1
        return f"スライド{op['slide']}「{sh.name}」の文字 {n} か所の色を {op.get('theme') or '#' + op.get('rgb', '').lstrip('#')} に"
    if kind == "theme-font":
        _, sh = find_shape(prs, op)
        n = 0
        for r in runs(sh):
            rpr = r._r.find("{%s}rPr" % NS_A)
            if rpr is None:
                continue
            for tag in ("latin", "ea", "cs"):
                f = rpr.find("{%s}%s" % (NS_A, tag))
                if f is not None and not (f.get("typeface") or "").startswith("+"):
                    rpr.remove(f)
                    n += 1
        return f"スライド{op['slide']}「{sh.name}」の直接指定のフォント {n} か所を外してテーマのフォントに"
    if kind == "lang":
        lang = op.get("lang", "ja-JP")
        n = 0
        for slide in prs.slides:
            for sh in walk(slide.shapes):
                for r in runs(sh):
                    if JA.search(r.text or ""):
                        rpr = r._r.get_or_add_rPr()
                        if rpr.get("lang") != lang:
                            rpr.set("lang", lang)
                            n += 1
        return f"日本語の文字 {n} か所に言語設定 {lang}"
    if kind == "replace-text":
        _, sh = find_shape(prs, op)
        n = 0
        for r in runs(sh):
            if op["old"] in r.text:
                r.text = r.text.replace(op["old"], op["new"])
                if JA.search(op["new"]):
                    r._r.get_or_add_rPr().set("lang", op.get("lang", "ja-JP"))
                n += 1
        if not n:
            raise FixError(f"図形 {sh.name} に「{op['old']}」がありません（1つのまとまりの中にある文字だけ置き換えられる）")
        return f"スライド{op['slide']}「{sh.name}」の「{op['old']}」を「{op['new']}」に"
    raise FixError(f"知らない op です: {kind}")


def main():
    ap = argparse.ArgumentParser(description="PowerPoint の指摘を直して別名で保存する")
    ap.add_argument("target")
    ap.add_argument("ops")
    ap.add_argument("--out")
    ap.add_argument("--pydir")
    a = ap.parse_args()
    if a.pydir:
        sys.path.insert(0, a.pydir)
    from pptx import Presentation

    out = a.out or re.sub(r"\.pptx$", "", a.target, flags=re.I) + ".dads-fixed.pptx"
    if os.path.abspath(out) == os.path.abspath(a.target):
        print("元のファイルには上書きしません。--out で別の名前を指定してください。", file=sys.stderr)
        sys.exit(2)
    ops = json.load(open(a.ops, encoding="utf-8"))
    prs = Presentation(a.target)
    masters_before = [m.part.blob for m in prs.slide_masters] + [l.part.blob for m in prs.slide_masters for l in m.slide_layouts]
    done, failed = [], []
    for op in ops:
        try:
            done.append(apply(prs, op))
        except (FixError, KeyError, IndexError) as e:
            failed.append(f"{json.dumps(op, ensure_ascii=False)} → {e}")
    if failed:
        print("直せない指示があったので、何も保存していません:", file=sys.stderr)
        for f in failed:
            print(f"  - {f}", file=sys.stderr)
        sys.exit(1)
    after = [m.part.blob for m in prs.slide_masters] + [l.part.blob for m in prs.slide_masters for l in m.slide_layouts]
    if after != masters_before:
        print("スライドマスター・レイアウトが変わってしまうため、保存しません。", file=sys.stderr)
        sys.exit(1)
    prs.save(out)
    print(json.dumps({"output": out, "applied": done}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
